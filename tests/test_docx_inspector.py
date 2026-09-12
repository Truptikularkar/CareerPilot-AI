import pytest
import docx
from pathlib import Path
from careerpilot.core.constants import FormattingRiskType, RiskSeverity
from careerpilot.ats.docx_inspector import DocxInspector
from careerpilot.generators.resume_docx import DocxResumeExporter
from careerpilot.graphs.resume_graph import generate_tailored_resume


def test_docx_inspector_clean_resume(tmp_path):
    eval_file = Path("data/jobs/evaluation/01_ai_data_engineer.txt")
    resume = generate_tailored_resume(eval_file)

    docx_path = tmp_path / "clean_resume.docx"
    DocxResumeExporter.export_docx(resume, docx_path)

    findings = DocxInspector.inspect_docx(docx_path)
    assert findings["file_exists"] is True
    assert findings["table_count"] == 0
    assert findings["header_has_content"] is False
    assert findings["footer_has_content"] is False
    # No high-severity formatting risks in our clean export
    high_risks = [r for r in findings["formatting_risks"] if r.severity in (RiskSeverity.HIGH, RiskSeverity.CRITICAL)]
    assert len(high_risks) == 0


def test_docx_inspector_table_trap_detection(tmp_path):
    # Create docx with a layout table
    doc = docx.Document()
    doc.add_paragraph("Trupti Kularkar")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python, SQL"

    trap_path = tmp_path / "table_trap_resume.docx"
    doc.save(str(trap_path))

    findings = DocxInspector.inspect_docx(trap_path)
    assert findings["table_count"] == 1
    assert any(r.risk_type == FormattingRiskType.TABLE_LAYOUT for r in findings["formatting_risks"])
    assert any(r.severity == RiskSeverity.HIGH for r in findings["formatting_risks"])


def test_docx_inspector_header_trap_detection(tmp_path):
    # Create docx with email hidden in header
    doc = docx.Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "kularkartrupti@gmail.com | +91 9834055766"
    doc.add_paragraph("Trupti Kularkar")

    trap_path = tmp_path / "header_trap_resume.docx"
    doc.save(str(trap_path))

    findings = DocxInspector.inspect_docx(trap_path)
    assert findings["header_has_content"] is True
    assert any(r.risk_type == FormattingRiskType.HEADER_FOOTER_CONTENT for r in findings["formatting_risks"])
