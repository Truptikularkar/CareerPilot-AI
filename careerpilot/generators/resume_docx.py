from pathlib import Path
from typing import Union, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from careerpilot.models.resume import TailoredResume
from careerpilot.models.artifact import ResumeArtifact
from careerpilot.generators.resume_pdf import PDFResumeExporter
from careerpilot.core.config import settings
from careerpilot.core.logging import get_logger

logger = get_logger(__name__)


class DocxResumeExporter:
    """
    Exports TailoredResume into a clean, ATS-compliant Microsoft Word (.docx) document.
    Uses single-column linear layout, standard fonts, and simple bullet formatting.
    """

    @classmethod
    def export_docx(
        cls,
        resume: TailoredResume,
        output_path: Optional[Union[str, Path]] = None,
        company: Optional[str] = None,
    ) -> ResumeArtifact:
        if output_path is not None:
            path = Path(output_path)
        else:
            path = Path(settings.BASE_DIR) / "data" / "generated" / "resumes" / resume.id / "resume.docx"

        path.parent.mkdir(parents=True, exist_ok=True)


        doc = docx.Document()

        # Set 0.6 inch standard margins
        for section in doc.sections:
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)

        # Set default font style
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(10.5)
        font.color.rgb = RGBColor(30, 30, 30)

        h = resume.header

        # 1. Header: Candidate Name
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(h.full_name)
        title_run.font.name = 'Calibri'
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(10, 37, 64)
        title_p.paragraph_format.space_after = Pt(2)

        # Contact Info Line
        contact_p = doc.add_paragraph()
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_parts = [p for p in [h.email, h.phone, h.location] if p]
        contact_run = contact_p.add_run(" | ".join(contact_parts))
        contact_run.font.size = Pt(9.5)
        contact_run.font.color.rgb = RGBColor(70, 70, 70)
        contact_p.paragraph_format.space_after = Pt(2)

        # Links Line
        links = [p for p in [h.linkedin_url, h.github_url] if p]
        if links:
            links_p = doc.add_paragraph()
            links_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            links_run = links_p.add_run(" | ".join(links))
            links_run.font.size = Pt(9.0)
            links_run.font.color.rgb = RGBColor(0, 102, 204)
            links_p.paragraph_format.space_after = Pt(8)

        def add_section_heading(title: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(title.upper())
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(10, 37, 64)

        # 2. Professional Summary
        add_section_heading("Professional Summary")
        summary_p = doc.add_paragraph()
        summary_p.paragraph_format.space_after = Pt(6)
        summary_run = summary_p.add_run(resume.summary.text)
        summary_run.font.size = Pt(10)

        # 3. Technical Skills
        add_section_heading("Technical Skills")
        for sc in resume.skills_categories:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(1.5)
            cat_run = p.add_run(f"{sc.category_name}: ")
            cat_run.font.bold = True
            cat_run.font.size = Pt(10)
            skills_run = p.add_run(", ".join(sc.skills))
            skills_run.font.size = Pt(10)

        # 4. Professional Experience
        add_section_heading("Professional Experience")
        for exp in resume.experiences:
            exp_header_p = doc.add_paragraph()
            exp_header_p.paragraph_format.space_before = Pt(4)
            exp_header_p.paragraph_format.space_after = Pt(1)
            
            title_run = exp_header_p.add_run(f"{exp.title} — {exp.company}")
            title_run.font.bold = True
            title_run.font.size = Pt(10.5)

            meta_p = doc.add_paragraph()
            meta_p.paragraph_format.space_after = Pt(3)
            loc_str = f" | {exp.location}" if exp.location else ""
            meta_run = meta_p.add_run(f"{exp.start_date} – {exp.end_date}{loc_str}")
            meta_run.font.italic = True
            meta_run.font.size = Pt(9.5)
            meta_run.font.color.rgb = RGBColor(90, 90, 90)

            for b in exp.bullets:
                bp = doc.add_paragraph(style='List Bullet')
                bp.paragraph_format.space_after = Pt(2)
                brun = bp.add_run(b)
                brun.font.size = Pt(10)

        # 5. Key Projects
        add_section_heading("Key Projects")
        for proj in resume.projects:
            proj_header_p = doc.add_paragraph()
            proj_header_p.paragraph_format.space_before = Pt(4)
            proj_header_p.paragraph_format.space_after = Pt(1)

            pname_run = proj_header_p.add_run(proj.name)
            pname_run.font.bold = True
            pname_run.font.size = Pt(10.5)

            if proj.technologies:
                tech_run = proj_header_p.add_run(f" (Technologies: {', '.join(proj.technologies)})")
                tech_run.font.italic = True
                tech_run.font.size = Pt(9.5)
                tech_run.font.color.rgb = RGBColor(90, 90, 90)

            for b in proj.bullets:
                bp = doc.add_paragraph(style='List Bullet')
                bp.paragraph_format.space_after = Pt(2)
                brun = bp.add_run(b)
                brun.font.size = Pt(10)

        # 6. Education
        add_section_heading("Education")
        for edu in resume.education:
            edu_p = doc.add_paragraph(style='List Bullet')
            edu_p.paragraph_format.space_after = Pt(2)
            deg_run = edu_p.add_run(f"{edu.degree} in {edu.field_of_study}")
            deg_run.font.bold = True
            deg_run.font.size = Pt(9.5)
            
            grad_txt = f" ({edu.graduation_year})" if edu.graduation_year else ""
            honors_txt = f" | {edu.honors}" if getattr(edu, "honors", None) else ""
            inst_run = edu_p.add_run(f" — {edu.institution}{grad_txt}{honors_txt}")
            inst_run.font.size = Pt(9.5)

        # 7. Certifications
        if resume.certifications:
            add_section_heading("Certifications")
            for cert in resume.certifications:
                cert_p = doc.add_paragraph(style='List Bullet')
                cert_p.paragraph_format.space_after = Pt(1.5)
                c_name = cert.name if hasattr(cert, "name") else str(cert)
                issuer_txt = f" — {cert.issuer}" if hasattr(cert, "issuer") and cert.issuer else ""
                year_txt = f" ({cert.year})" if hasattr(cert, "year") and cert.year else ""
                crun = cert_p.add_run(f"{c_name}{issuer_txt}{year_txt}")
                crun.font.size = Pt(9.5)


        doc.save(str(path))
        logger.info("Successfully exported ATS DOCX resume to: %s", path)

        role_name = "Engineer"
        if hasattr(resume, "strategy") and resume.strategy and hasattr(resume.strategy, "strategy_type"):
            role_name = resume.strategy.strategy_type.value

        safe_pdf_name = PDFResumeExporter.get_safe_filename(
            candidate_name=resume.header.full_name,
            role_title=role_name,
            company=company,
        )
        safe_filename = safe_pdf_name.replace(".pdf", ".docx")

        return ResumeArtifact(
            artifact_type="docx",
            file_name=safe_filename,
            file_path=str(path),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            resume_id=resume.id,
            is_valid=True,
        )

